import csv
import io
import re
import zipfile
from pathlib import Path

import duckdb
import pandas as pd
from openpyxl import load_workbook

from app.config import Settings
from app.errors import AppError
from app.serialization import json_value


def invalid(message: str, **details):
    return AppError(422, "invalid_file", message, **details)


def validate_headers(raw: list, settings: Settings) -> list[str]:
    if len(raw) > settings.max_dataset_columns:
        raise invalid(f"列数超过限制 {settings.max_dataset_columns}。")
    headers = [str(value).strip() if value is not None else "" for value in raw]
    if not headers or any(not name for name in headers):
        raise invalid("列名不能为空，请补全表头。")
    if len({name.casefold() for name in headers}) != len(headers):
        raise invalid("存在重复列名（忽略大小写和首尾空白），请重命名。")
    if any(len(name) > 128 or any(ord(c) < 32 for c in name) for name in headers):
        raise invalid("列名不能包含控制字符，且长度不能超过 128。")
    return headers


def read_csv(content: bytes, settings: Settings) -> pd.DataFrame:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise invalid("CSV 必须使用 UTF-8 编码。") from exc
    if "\x00" in text or any(ord(c) < 32 and c not in "\r\n\t" for c in text):
        raise invalid("CSV 包含二进制数据或无效控制字符。")
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    try:
        reader = csv.reader(io.StringIO(text, newline=""), dialect=dialect, strict=True)
        headers = validate_headers(next(reader, []), settings)
        rows = []
        for row in reader:
            if not row or all(not value.strip() for value in row):
                continue
            if len(row) != len(headers):
                raise invalid(f"CSV 第 {reader.line_num} 行的字段数与表头不一致。")
            rows.append(row)
            if len(rows) > settings.max_dataset_rows:
                raise invalid(f"行数超过限制 {settings.max_dataset_rows}。")
    except csv.Error as exc:
        raise invalid("CSV 格式无效或单个字段过长。") from exc
    return pd.DataFrame(rows, columns=headers)


def read_xlsx(content: bytes, sheet: str | None, settings: Settings):
    if not zipfile.is_zipfile(io.BytesIO(content)):
        raise invalid("文件内容不是有效的 XLSX 工作簿。")
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            members = archive.infolist()
            names = {info.filename for info in members}
            if not {"[Content_Types].xml", "xl/workbook.xml"}.issubset(names):
                raise invalid("缺少 XLSX 必需的文件结构。")
            if len(members) > 1000 or sum(info.file_size for info in members) > 50 * 1024**2:
                raise invalid("XLSX 解压后超过资源限制（50 MiB / 1000 个条目）。")
            if any("vbaProject" in name or name.startswith("xl/externalLinks/") for name in names):
                raise invalid("不接受宏或外部工作簿链接。")
        workbook = load_workbook(
            io.BytesIO(content), read_only=True, data_only=False, keep_links=False
        )
        try:
            sheets = workbook.sheetnames
            if sheet is None and len(sheets) != 1:
                raise AppError(422, "sheet_required", "请明确选择工作表。", sheets=sheets)
            selected = sheet if sheet is not None else sheets[0]
            if selected not in sheets:
                raise invalid("指定的工作表不存在。", sheets=sheets)
            worksheet = workbook[selected]
            if (worksheet.max_row or 0) > settings.max_dataset_rows + 1:
                raise invalid(f"工作表行数超过限制 {settings.max_dataset_rows}。")
            if (worksheet.max_column or 0) > settings.max_dataset_columns:
                raise invalid(f"工作表列数超过限制 {settings.max_dataset_columns}。")
            iterator = worksheet.iter_rows(
                max_row=min(worksheet.max_row or 1, settings.max_dataset_rows + 2),
                max_col=min(worksheet.max_column or 1, settings.max_dataset_columns + 1),
            )
            header_cells = next(iterator, ())
            if any(cell.data_type == "f" for cell in header_cells):
                raise invalid("不接受公式单元格，请将公式转换为值后上传。")
            headers = validate_headers([cell.value for cell in header_cells], settings)
            rows = []
            for cells in iterator:
                if any(cell.data_type == "f" for cell in cells):
                    raise invalid("不接受公式单元格，请将公式转换为值后上传。")
                values = [cell.value for cell in cells]
                if all(value is None for value in values):
                    continue
                rows.append(values)
                if len(rows) > settings.max_dataset_rows:
                    raise invalid(f"行数超过限制 {settings.max_dataset_rows}。")
            return pd.DataFrame(rows, columns=headers), selected
        finally:
            workbook.close()
    except AppError:
        raise
    except Exception as exc:
        raise invalid("无法解析 XLSX，请确认文件未损坏、未加密。") from exc


def normalize(frame: pd.DataFrame) -> pd.DataFrame:
    """Conservative whole-column inference; mixed types remain text. No eval/formula execution."""
    for name in frame:
        series = frame[name].map(lambda v: None if isinstance(v, str) and not v.strip() else v)
        nonnull = series.dropna()
        if nonnull.empty:
            frame[name] = series.astype("string")
            continue
        strings = nonnull.astype(str)
        # Leading zero codes (e.g. 00123) are identifiers, not numbers.
        leading_zero = strings.str.match(r"^[+-]?0\d+$").any()
        numeric = pd.to_numeric(nonnull, errors="coerce")
        if not leading_zero and numeric.notna().all():
            frame[name] = pd.to_numeric(series, errors="coerce").replace(
                [float("inf"), float("-inf")], pd.NA
            )
        elif strings.str.lower().isin(["true", "false"]).all():
            frame[name] = series.map(
                lambda v: None if pd.isna(v) else str(v).lower() == "true"
            ).astype("boolean")
        elif strings.map(lambda v: bool(re.fullmatch(r"\d{4}-\d{2}-\d{2}", v))).all():
            dates = pd.to_datetime(series, errors="coerce")
            frame[name] = dates if dates.notna().sum() == len(nonnull) else series.astype("string")
        elif pd.api.types.is_datetime64_any_dtype(series):
            frame[name] = series
        else:
            frame[name] = series.astype("string")
    return frame.convert_dtypes()


def import_dataset(content: bytes, suffix: str, sheet: str | None, path: Path, settings: Settings):
    if not content:
        raise invalid("文件为空。")
    if len(content) > settings.max_upload_bytes:
        raise AppError(413, "file_too_large", f"文件大小超过 {settings.max_upload_bytes} 字节。")
    if suffix == ".csv":
        if sheet:
            raise invalid("CSV 不支持工作表参数。")
        frame, selected = read_csv(content, settings), None
    elif suffix == ".xlsx":
        frame, selected = read_xlsx(content, sheet, settings)
    else:
        raise AppError(415, "unsupported_format", "仅支持 .csv 和 .xlsx。")
    if frame.empty:
        raise invalid("没有数据行，请保留表头并至少提供一行数据。")
    frame = normalize(frame)
    with duckdb.connect(str(path), config={"threads": 2, "memory_limit": "256MB"}) as connection:
        connection.register("incoming", frame)
        connection.execute('CREATE TABLE "data" AS SELECT * FROM incoming')
        schema = connection.execute('DESCRIBE "data"').fetchall()
    fields = []
    for name, dtype, *_ in schema:
        series = frame[name]
        missing = int(series.isna().sum())
        stats = {}
        if pd.api.types.is_numeric_dtype(series) and not pd.api.types.is_bool_dtype(series):
            for key, value in series.describe().items():
                stats[key] = json_value(value)
        fields.append(
            {
                "name": name,
                "type": dtype,
                "missing": missing,
                "missing_ratio": missing / len(frame),
                "unique": int(series.nunique()),
                "statistics": stats,
            }
        )
    preview = [
        {
            name: (value[:500] if isinstance(value, str) else value)
            for name, raw in record.items()
            for value in [json_value(raw)]
        }
        for record in frame.head(20).to_dict(orient="records")
    ]
    return {
        "table_name": "data",
        "row_count": len(frame),
        "column_count": len(frame.columns),
        "fields": fields,
        "preview": preview,
        "notes": ["预览最多 20 行，文本最多 500 字符；类型由整列保守推断，空白视为缺失。"],
    }, selected
